from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


SHARES_COLUMN = 11  # column K: Part II holdings
FIRST_ROW = 21      # row 21 is the risk-free line, stocks start at row 22


def export_to_excel(weights: pd.Series, capital: float, template: Path, output: Path):
    """
    Purpose:
        Convert weights into whole shares and write them to the hand-in sheet

    Inputs:
        weights     final portfolio weights indexed by PERMNO (excluding risk-free)
        capital     total budget
        template    blank hand-in sheet
        output      path of the completed sheet

    Returns:
        DataFrame with shares, amounts and realized weights per stock
    """
    df = pd.read_excel(template, header=FIRST_ROW - 2, usecols='A:F')
    df = df.loc[df['PERMNO'].notna()].copy()
    df['PERMNO'] = df['PERMNO'].astype(int)
    price = df['Close price 31-12-2025']

    missing = weights.index.difference(df['PERMNO'])
    if len(missing) > 0:
        raise ValueError(f'PERMNOs not in the hand-in sheet: {missing.tolist()}')

    df['target_weight'] = df['PERMNO'].map(weights).fillna(0)
    df['shares'] = np.floor(capital * df['target_weight'] / price).astype(int)
    df['investment'] = df['shares'] * price
    df['weight'] = df['investment'] / capital
    rf_investment = capital - df['investment'].sum()

    wb = load_workbook(template)
    ws = wb.worksheets[0]
    for i, shares in df['shares'].items():
        ws.cell(row=int(i) + FIRST_ROW, column=SHARES_COLUMN, value=int(shares))
    ws.cell(row=FIRST_ROW, column=SHARES_COLUMN, value=float(rf_investment))
    ws['C12'] = capital

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return df
