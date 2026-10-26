"""
Final Part III portfolio for January 2026 and the hand-in sheet.

Expected returns come from the Random Forest forecast for Jan 2026 (outputs/jan_2026_forecast.csv). The
model forecasts the excess return over the T-bill, so the latest T-bill return (Dec 2025) is added to get an
expected total return. The optimizer then compares that with the assignment's risk-free rate of 1% per year.
The covariance uses the 2023-2025 daily returns, same window as Part II.
"""

from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from lib.covariance_model import load_daily_panel, estimate_covariance
from lib.portfolio import MeanVariancePortfolio

CAPITAL = 10_000_000
RF_MONTHLY = 1.01 ** (1 / 12) - 1     # assignment risk-free rate: 1% per year
CAP = 0.10
GAMMA = 4
FINAL_MODEL = 'pred_rf'
COV_START, COV_END = '2023-01-01', '2025-12-31'
# group's copy of the hand-in sheet (Parts I and II filled in). the export matches stocks by PERMNO,
# so the official blank template works here too
TEMPLATE = 'portfolio_template_part2.xlsx'

SHARES_COLUMN = 14   # column N: Part III holdings
FIRST_ROW = 21       # row 21 is the risk-free line, stocks start at row 22


def sheet_stocks(template):
    """
    PERMNOs and 31-12-2025 prices in the order of the hand-in sheet, with the Excel row of each stock
    """
    df = pd.read_excel(template, header = FIRST_ROW - 2, usecols = 'A:F')
    # index 0 is the risk-free line (row 21), so index i sits in Excel row i + 21
    df['Row'] = df.index + FIRST_ROW
    df = df.loc[df['PERMNO'].notna()].copy()
    df['PERMNO'] = df['PERMNO'].astype(int)
    return df


def firm_list_permnos(data_dir: Path, template: Path):
    """
    PERMNOs we are allowed to hold: stocks in the hand-in sheet whose ticker is on the provided firm list.
    This drops rows that are in the sheet but not on the list (AMR and METV in our group's copy).
    """
    tickers = set(pd.read_csv(data_dir / 'assignment_tickers_2025.txt', header = None)[0].str.strip())
    sheet = sheet_stocks(template)
    return pd.Index(sheet.loc[sheet['Ticker Symbol'].isin(tickers), 'PERMNO'].unique())


def build_final_portfolio(outputs_dir: Path, data_dir: Path, template: Path, model = FINAL_MODEL):
    """
    Mean-variance portfolio for Jan 2026 from one model's forecast

    Returns:
        MeanVariancePortfolio object
    """
    forecast = pd.read_csv(outputs_dir / 'jan_2026_forecast.csv')
    stocks = pd.read_csv(outputs_dir / 'stock_monthly.csv', usecols = ['Month', 'Rf'], parse_dates = ['Month'])
    tbill = stocks.loc[stocks['Month'] == '2025-12-01', 'Rf'].iloc[0]

    allowed = firm_list_permnos(data_dir, template)
    forecast = forecast.loc[forecast['PERMNO'].isin(allowed)].set_index('PERMNO')
    mu = forecast[model] + tbill

    panel = load_daily_panel(data_dir)
    cov, shrinkage = estimate_covariance(panel, mu.index, COV_START, COV_END)
    print(f'Stocks with a forecast: {len(mu)}, with full 2023-2025 daily data: {len(cov)}, shrinkage {shrinkage:.3f}')
    return MeanVariancePortfolio(mu.reindex(cov.index), cov, rf = RF_MONTHLY, cap = CAP, gamma = GAMMA)


def export_to_excel(weights: pd.Series, capital: float, template: Path, output: Path):
    """
    Turn weights into whole shares at the 31-12-2025 prices and write them in column N.
    Whatever is left after rounding down goes to the risk-free line.

    Returns:
        DataFrame with shares, amount invested and realized weight per stock
    """
    df = sheet_stocks(template)
    price = df['Close price 31-12-2025']

    missing = weights.index.difference(df['PERMNO'])
    if len(missing) > 0:
        raise ValueError(f'PERMNOs not in the hand-in sheet: {missing.tolist()}')

    # match on PERMNO, not on row position. if a PERMNO appears twice (HST in the official file),
    # only its first row gets the position so it is not bought twice
    df['target_weight'] = df['PERMNO'].map(weights).fillna(0).where(~df['PERMNO'].duplicated(), 0)
    df['shares'] = np.floor(capital * df['target_weight'] / price).astype(int)
    df['investment'] = df['shares'] * price
    df['weight'] = df['investment'] / capital
    rf_investment = capital - df['investment'].sum()
    if rf_investment < 0 or (df['shares'] < 0).any():
        raise ValueError('Budget exceeded or negative position')

    wb = load_workbook(template)
    ws = wb.worksheets[0]
    for row, shares in zip(df['Row'], df['shares']):
        ws.cell(row = int(row), column = SHARES_COLUMN, value = int(shares))
    ws.cell(row = FIRST_ROW, column = SHARES_COLUMN, value = float(rf_investment))

    output.parent.mkdir(parents = True, exist_ok = True)
    wb.save(output)
    return df, rf_investment


def run_final_portfolio(outputs_dir: Path, data_dir: Path):
    template = data_dir / TEMPLATE
    portfolio = build_final_portfolio(outputs_dir, data_dir, template)
    summary = portfolio.summary()
    summary.to_csv(outputs_dir / 'portfolio_summary.csv')

    holdings, rf_investment = export_to_excel(portfolio.weights, CAPITAL, template, outputs_dir / 'Hand_in_sheet.xlsx')
    held = holdings.loc[holdings['shares'] > 0, ['PERMNO', 'Ticker Symbol', 'Company Name', 'Close price 31-12-2025',
                                                 'target_weight', 'shares', 'investment', 'weight']]
    held.sort_values('weight', ascending = False).to_csv(outputs_dir / 'holdings.csv', index = False)

    print(summary.round(4).to_string())
    print(f'Invested in stocks: {holdings["investment"].sum():,.2f}')
    print(f'Invested at risk-free: {rf_investment:,.2f}')
    return summary, held


if __name__ == '__main__':
    root = Path(__file__).resolve().parent.parent
    run_final_portfolio(root / 'outputs', root / 'data')
